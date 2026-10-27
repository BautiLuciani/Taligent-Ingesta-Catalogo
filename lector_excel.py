"""Lector de Excel: lee la tabla de productos de una hoja, con la fuente {archivo, hoja, fila} de cada uno.

Semana 2, paso 6C. Sirve para todos los Excel del proyecto, aunque cada uno sea distinto (encabezado en la
fila 1, 4 o 5; "Código", "Item", "COD. RIV" o "Item #"; títulos de sección en el medio). No busca un texto
fijo: busca la primera fila cuya primera celda es un código de producto, y el encabezado es la fila
anterior más cercana con 2 o más celdas. Los valores quedan TAL CUAL (las cuentas las hace el servidor).

Trampa de Rivera: sus títulos "== CAPILAR ==" empiezan con "=", así que Excel los guarda como fórmulas sin
valor, y leyendo solo los valores vendrían vacíos. Por eso una fórmula sin valor guardado se toma como texto.
"""
from pathlib import Path

import openpyxl

from cuentas import leer_fecha
from extractor_pdf import codigo_al_inicio


def es_codigo(valor) -> bool:
    """¿La celda es un código de producto entero? "AN-1124" sí; "== CAPILAR ==", "Código" o 7790670659075, no."""
    if not isinstance(valor, str):
        return False
    texto = valor.strip()
    return codigo_al_inicio(texto) == texto


def leer_tabla_excel(ruta: Path, hoja: str | None = None) -> dict:
    """Lee la tabla de productos de una hoja (por defecto, la primera).

    Devuelve un diccionario con:
      - filas: una por código, con su fuente {archivo, hoja, fila} (la fila que muestra Excel) y su "original":
        {nombre de columna: valor tal cual} (un número queda número, una celda vacía queda None).
      - vigencia: la fecha de una celda de arriba del encabezado que diga "vigencia", con su fuente; o None.
      - ignoradas: las filas con algo escrito que no son un producto (títulos de sección): nada desaparece en silencio.
      - problemas: lo que impidió leer algo (sin filas de productos, códigos repetidos, dos vigencias).
    """
    libro = openpyxl.load_workbook(ruta, data_only=True)       # los valores que se ven
    formulas = openpyxl.load_workbook(ruta, data_only=False)   # el texto de cada celda, fórmulas incluidas
    planilla = libro[hoja] if hoja else libro.worksheets[0]
    planilla_formulas = formulas[planilla.title]
    archivo, nombre_hoja = Path(ruta).name, planilla.title

    def valor_de(celda):
        """El valor de la celda; si es una fórmula sin valor guardado ("== CAPILAR =="), su texto tal cual."""
        texto = planilla_formulas[celda.coordinate]
        return texto.value if celda.value is None and texto.data_type == "f" else celda.value

    renglones = [(celdas[0].row, [valor_de(celda) for celda in celdas]) for celdas in planilla.iter_rows()]
    resultado = {"archivo": archivo, "hoja": nombre_hoja, "fila_encabezado": None, "columnas": [],
                 "vigencia": None, "filas": [], "ignoradas": [], "problemas": []}

    def no_vacias(valores):
        return [v for v in valores if v is not None and str(v).strip()]

    # 1) La primera fila de datos: la primera cuya primera celda con algo escrito es un código
    primera = next((i for i, (_, valores) in enumerate(renglones)
                    if no_vacias(valores) and es_codigo(no_vacias(valores)[0])), None)
    if primera is None:
        resultado["problemas"].append(f"La hoja {nombre_hoja!r} no tiene ninguna fila que empiece con un código.")
        return resultado

    # 2) El encabezado: la fila anterior más cercana con 2 o más celdas (saltea títulos de sección)
    encabezado = next((i for i in range(primera - 1, -1, -1) if len(no_vacias(renglones[i][1])) >= 2), None)
    if encabezado is None:
        resultado["problemas"].append(f"La hoja {nombre_hoja!r} no tiene encabezado arriba de los productos.")
        return resultado
    fila_encabezado, nombres = renglones[encabezado]
    columnas = [str(nombre).strip() if nombre is not None else None for nombre in nombres]
    resultado["fila_encabezado"] = fila_encabezado
    resultado["columnas"] = [c for c in columnas if c]
    col_codigo = next(j for j, v in enumerate(renglones[primera][1]) if v is not None and str(v).strip())

    # 3) La vigencia: una celda de ARRIBA del encabezado que diga "vigencia" (ej. "Vigencia desde: 15/10/2026")
    vigencias = {}
    for fila_excel, valores in renglones[:encabezado]:
        for valor in no_vacias(valores):
            texto = str(valor).strip()
            fecha = leer_fecha(texto) if "vigencia" in texto.lower() else None
            if fecha is not None:
                vigencias.setdefault(fecha, {"fecha": fecha.isoformat(), "texto": texto,
                                             "fuente": {"archivo": archivo, "hoja": nombre_hoja, "fila": fila_excel}})
    if len(vigencias) == 1:
        resultado["vigencia"] = next(iter(vigencias.values()))
    elif len(vigencias) > 1:
        fechas = ", ".join(sorted(v["fecha"] for v in vigencias.values()))
        resultado["problemas"].append(f"La hoja cita más de una vigencia ({fechas}): no se elige una.")

    # 4) Los productos: cada fila de abajo cuya celda de código es un código; lo demás escrito, a "ignoradas"
    por_codigo, repetidos = {}, set()
    for fila_excel, valores in renglones[encabezado + 1:]:
        celda = valores[col_codigo] if col_codigo < len(valores) else None
        if not es_codigo(celda):
            if no_vacias(valores):
                resultado["ignoradas"].append({"fila": fila_excel, "texto": " ".join(map(str, no_vacias(valores)))})
            continue
        codigo = celda.strip()
        original = {nombre: valores[j] if j < len(valores) else None
                    for j, nombre in enumerate(columnas) if nombre and j != col_codigo}
        nueva = {"codigo": codigo, "fuente": {"archivo": archivo, "hoja": nombre_hoja, "fila": fila_excel},
                 "original": original}
        if codigo in por_codigo and por_codigo[codigo]["original"] != original:
            repetidos.add(codigo)
        por_codigo.setdefault(codigo, nueva)  # repetido idéntico: queda el primero

    resultado["problemas"] += [f"{codigo}: aparece dos veces con datos distintos: no se elige una." for codigo in sorted(repetidos)]
    resultado["filas"] = [fila for codigo, fila in por_codigo.items() if codigo not in repetidos]
    return resultado